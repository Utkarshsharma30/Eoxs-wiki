"""One-off import: merges EOXS's messy multi-sheet employee spreadsheet into
one canonical row per person for the new `employees` table.

Not a recurring ingestion source (unlike everything in ingestion/) -- this
runs once (or occasionally, by hand, pointed at a fresh export), same
category as loaders/'s existing one-time historical-data loaders.

Deliberately reads only 3 of the workbook's 5 sheets, and only a subset of
their columns -- per an explicit decision: no LinkedIn URLs (SHEET1), no
Google-Review contact rows (Google Review), no personal phone numbers, and
no plaintext account passwords -- the "Employee name " sheet has a Password
column for personal Gmail logins (same value, 'Eoxs12345!', for nearly
everyone); it is never read here, on purpose, regardless of how this table
gets used later.

  "Active Employee Name & Email ID"  -> name, department (fallback), official_email (fallback)
  "Employee name "                   -> name, department (fallback), employment_type, official_email (fallback)
  "SHEET2"                           -> name, department (preferred), role_title, manager (TL column),
                                         official_email (preferred), date_of_joining

SHEET2's header row has a genuine duplicate: two columns are both literally
labeled "Department" (one holds a job title, the next one the real
department) -- dict(zip(header, row)) silently collapses that to the LAST
column's value, losing the job title entirely. Handled below by reading
SHEET2 positionally instead of by header name.

People are merged across sheets by a normalized name key (lowercased,
whitespace-collapsed, trailing-period-stripped) -- the sheets don't share a
stable id, and several names have inconsistent trailing spaces
("Arpita " vs "Arpita") or casing across sheets.

Usage:
  python -m loaders.import_employees_from_xlsx <path-to-xlsx>            # preview only, no DB writes
  python -m loaders.import_employees_from_xlsx <path-to-xlsx> --commit   # actually create_employee() each row
                                                                          #   (writes to PGDATABASE from .env --
                                                                          #    e.g. PGDATABASE=eoxs_wiki_staging
                                                                          #    to test first, same convention as
                                                                          #    schema/run_migrations.py)
Every row is created with status='active' (this workbook is a current-roster
snapshot, not a historical record -- no one in it is marked as having left).
Rows whose normalized name already exists in the `employees` table (by
official_email, or by name if no email) are skipped, not duplicated -- safe
to re-run against an updated export.
"""
import re
import sys

import openpyxl

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
OFFICIAL_DOMAINS = ("@eoxs.com", "@eoxsteam.com")


def normalize_name(name):
    if not name:
        return ""
    name = re.sub(r"\s+", " ", str(name)).strip()
    return name.rstrip(".").strip().lower()


def display_name(name):
    return re.sub(r"\s+", " ", str(name)).strip()


def best_email(*raw_fields):
    """Pulls every email-looking substring out of the given fields (some
    cells hold several, comma/space separated) and prefers an official
    eoxs.com/eoxsteam.com address over a personal one."""
    found = []
    for field in raw_fields:
        if not field:
            continue
        found.extend(EMAIL_RE.findall(str(field)))
    if not found:
        return None
    for e in found:
        if any(d in e.lower() for d in OFFICIAL_DOMAINS):
            return e.strip().lower()
    return found[0].strip().lower()


def _dict_rows(ws, header_row_idx=1):
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[header_row_idx - 1]]
    out = []
    for row in rows[header_row_idx:]:
        if not any(c is not None for c in row):
            continue
        out.append(dict(zip(header, row)))
    return out


def _positional_rows(ws, header_row_idx=1):
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for row in rows[header_row_idx:]:
        if not any(c is not None for c in row):
            continue
        out.append(row)
    return out


def merge(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    active_rows = _dict_rows(wb["Active Employee Name & Email ID"])
    empname_rows = _dict_rows(wb["Employee name "])
    # SHEET2 columns, positional (0-indexed): 0=Sr.No 1=Name 2=role_title
    # 3=department 4=TL/manager 5=EOXS_PORTAL 6=Contact 7=Email 8=Date Of Joining
    sheet2_rows = _positional_rows(wb["SHEET2"])

    people = {}  # normalized_name -> record

    def get_or_create(name):
        key = normalize_name(name)
        if not key:
            return None
        if key not in people:
            people[key] = {
                "full_name": display_name(name), "department": None, "role_title": None,
                "employment_type": None, "official_email": None, "manager": None,
                "date_of_joining": None, "_sources": set(),
            }
        return people[key]

    for r in active_rows:
        p = get_or_create(r.get("Name"))
        if not p:
            continue
        p["_sources"].add("active_list")
        p["department"] = p["department"] or (str(r.get("Department") or "").strip() or None)
        p["official_email"] = p["official_email"] or best_email(r.get("Login/Account Used (Official/Personal)"))

    for r in empname_rows:
        p = get_or_create(r.get("Name"))
        if not p:
            continue
        p["_sources"].add("employee_name_sheet")
        p["department"] = p["department"] or (str(r.get("Department") or "").strip() or None)
        p["employment_type"] = p["employment_type"] or (str(r.get("Type Of Employee") or "").strip() or None)
        p["official_email"] = p["official_email"] or best_email(r.get("Email"))

    for row in sheet2_rows:
        name, role_title, department, manager = row[1], row[2], row[3], row[4]
        email, date_of_joining = row[7], row[8]
        p = get_or_create(name)
        if not p:
            continue
        p["_sources"].add("sheet2")
        # SHEET2's own data is the most specific/current for these three --
        # preferred over the other two sheets' values, not just a fallback.
        if role_title:
            p["role_title"] = str(role_title).strip()
        if department:
            p["department"] = str(department).strip()
        if manager:
            p["manager"] = str(manager).strip()
        sheet2_email = best_email(email)
        if sheet2_email:
            p["official_email"] = sheet2_email
        if date_of_joining:
            try:
                p["date_of_joining"] = date_of_joining.strftime("%Y-%m-%d")
            except AttributeError:
                pass

    people = _merge_by_email(people)

    for p in people.values():
        p["source_count"] = len(p.pop("_sources"))

    return people, len(active_rows), len(empname_rows), len(sheet2_rows)


def _merge_by_email(people):
    """Second dedup pass: two DIFFERENT name spellings sharing the same
    official_email are almost certainly the same person (confirmed real
    case: 'Dhrup' in two sheets vs 'Dhrup Kumar' in SHEET2, identical
    dhrup@eoxsteam.com in both) -- name-key dedup alone can't catch this.
    Keeps whichever record has more populated fields (ties broken by
    longer name, since the fuller SHEET2-style name is usually the more
    complete/real one), filling in any still-blank field from the other."""
    by_email = {}
    for key, p in people.items():
        if not p["official_email"]:
            by_email.setdefault(("__no_email__", key), []).append(p)
            continue
        by_email.setdefault(p["official_email"], []).append(p)

    merged = {}
    for _, group in by_email.items():
        if len(group) == 1:
            winner = group[0]
        else:
            def completeness(p):
                return sum(1 for f in ("department", "role_title", "employment_type", "manager", "date_of_joining") if p[f])
            group.sort(key=lambda p: (completeness(p), len(p["full_name"])), reverse=True)
            winner = group[0]
            for other in group[1:]:
                for field in ("department", "role_title", "employment_type", "manager", "date_of_joining"):
                    winner[field] = winner[field] or other[field]
                winner["_sources"] |= other["_sources"]
        merged[normalize_name(winner["full_name"])] = winner
    return merged


def _trunc(s, n):
    s = s or ""
    return (s[: n - 1] + "…") if len(s) > n else s


def print_preview(people):
    rows = sorted(people.values(), key=lambda p: p["full_name"].lower())
    cols = [("Name", "full_name", 24), ("Department", "department", 20), ("Role/Title", "role_title", 28),
            ("Type", "employment_type", 11), ("Email", "official_email", 32), ("Manager", "manager", 14),
            ("Joined", "date_of_joining", 11)]
    print("  ".join(f"{label:<{w}}" for label, _, w in cols))
    print("-" * (sum(w for *_, w in cols) + 2 * (len(cols) - 1)))
    for p in rows:
        print("  ".join(f"{_trunc(p[field], w):<{w}}" for _, field, w in cols))
    print("-" * (sum(w for *_, w in cols) + 2 * (len(cols) - 1)))
    n = len(rows)
    print(f"\n{n} unique people merged.")
    print(f"  have department:      {sum(1 for p in rows if p['department'])}/{n}")
    print(f"  have role_title:      {sum(1 for p in rows if p['role_title'])}/{n}")
    print(f"  have employment_type: {sum(1 for p in rows if p['employment_type'])}/{n}")
    print(f"  have official_email:  {sum(1 for p in rows if p['official_email'])}/{n}")
    print(f"  have manager:         {sum(1 for p in rows if p['manager'])}/{n}")
    print(f"  have date_of_joining: {sum(1 for p in rows if p['date_of_joining'])}/{n}")
    print(f"  appear in only 1 sheet (least-verified): {sum(1 for p in rows if p['source_count'] == 1)}/{n}")


def commit(people):
    from mcp_server import employees as emp
    from mcp_server.db import query

    existing_emails = {row["official_email"] for row in query("SELECT official_email FROM employees WHERE official_email IS NOT NULL")}
    existing_names = {row["full_name"].strip().lower() for row in query("SELECT full_name FROM employees")}

    created, skipped = 0, 0
    for p in sorted(people.values(), key=lambda p: p["full_name"].lower()):
        if (p["official_email"] and p["official_email"] in existing_emails) or \
           (not p["official_email"] and p["full_name"].strip().lower() in existing_names):
            skipped += 1
            continue
        emp.create_employee(
            full_name=p["full_name"], department=p["department"] or "", role_title=p["role_title"] or "",
            employment_type=p["employment_type"] or "", official_email=p["official_email"] or "",
            manager=p["manager"] or "", date_of_joining=p["date_of_joining"] or "",
            notes="", changed_by="import_script",
        )
        created += 1
    print(f"created {created}, skipped {skipped} already-present")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: python -m loaders.import_employees_from_xlsx <path-to-xlsx> [--commit]")
        sys.exit(1)
    people, n_active, n_empname, n_sheet2 = merge(sys.argv[1])
    print(f"source rows: active_list={n_active} employee_name_sheet={n_empname} sheet2={n_sheet2}\n")
    print_preview(people)
    if "--commit" in sys.argv:
        print()
        commit(people)
