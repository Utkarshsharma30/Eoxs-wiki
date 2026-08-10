"""Text extraction for email/ticket attachments -- shared by the live
fetchers (gmail_fetcher.py, zoho_fetcher.py, called at ingestion time) and
ingestion/backfill_attachments.py (called against already-ingested rows).

Scope: PDF, DOCX, XLSX, TXT, CSV only. Images, .ics invites, archives, and
other binary formats are deliberately out of scope (no OCR) -- callers
still get mimetype/source_attachment_id filled from free API metadata for
those, just no extracted_text.
"""
import io
import os

import docx
import openpyxl
import pypdf

MAX_EXTRACT_BYTES = 20 * 1024 * 1024


def _extract_pdf_text(data):
    reader = pypdf.PdfReader(io.BytesIO(data))
    parts = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(p for p in parts if p).strip()
    return text or None


def _extract_docx_text(data):
    document = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs if p.text).strip()
    return text or None


def _extract_xlsx_text(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append("\t".join(cells))
    return "\n".join(lines).strip() or None


def _extract_plain_text(data):
    for encoding in ("utf-8", "latin-1"):
        try:
            return data.decode(encoding).strip() or None
        except UnicodeDecodeError:
            continue
    return None


EXTRACTORS = {
    ".pdf": _extract_pdf_text,
    ".docx": _extract_docx_text,
    ".xlsx": _extract_xlsx_text,
    ".txt": _extract_plain_text,
    ".csv": _extract_plain_text,
}


def is_extractable(filename):
    return os.path.splitext(filename)[1].lower() in EXTRACTORS


def extract_text(data, filename):
    """Returns extracted text, or None for unsupported types, oversized
    data, or a genuinely empty extraction result. Postgres text columns
    reject NUL bytes, which some PDF extractions produce -- stripped here
    so a real bug elsewhere can't lose an otherwise-good extraction."""
    extractor = EXTRACTORS.get(os.path.splitext(filename)[1].lower())
    if not extractor or len(data) > MAX_EXTRACT_BYTES:
        return None
    text = extractor(data)
    return text.replace("\x00", "") if text else text
